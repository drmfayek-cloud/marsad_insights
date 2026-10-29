"""End-to-end checks on synthetic data. Run: python -m unittest discover -s tests"""
import datetime as dt
import io
import os
import re
import shutil
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
TMP = tempfile.mkdtemp(prefix="insights-test-")
os.environ.update(MARSAD_DATA=TMP, MARSAD_DEMO="0", MARSAD_HTTPS="0", MARSAD_REQUIRE_MFA="0",
                  MARSAD_ADMIN_EMAIL="admin@test.local", MARSAD_ADMIN_PASSWORD="First!Pass2026x")

import app as A  # noqa: E402
from insights import auth, config, db, kpi, store  # noqa: E402

NEW_PW = "Changed!Pass2026x"


def tok(html):
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


def daily_register():
    """A minimal workbook with the same sheets and headers as the MARSAD Daily Record & Action Tracker."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Intake"
    ws.append(["Daily Intake"]); ws.append([]); ws.append([])
    ws.append(["Date received", "Received by", "Source", "TPA", "Volume received", "Reviewed", "Outstanding"])
    for d in range(1, 11):
        ws.append([dt.date(2026, 9, d), "Someone", "Approval", "GlobeMed", 100, 100, 0])
    ws.append([dt.date(2026, 10, 1), "Someone", "Approval", "Nextcare", 50, 40, 10])
    rl = wb.create_sheet("Daily Review Log")
    rl.append(["x"]); rl.append([]); rl.append([])
    rl.append(["Date", "Reviewer", "Source", "TPA", "Records reviewed", "Confirmed", "Suspected", "Flag rate %",
               "Declared total", "Logged in Daily Log"])
    rl.append([dt.date(2026, 9, 1), "Dr Example", "Approval", "GlobeMed", 100, 3, 0, 0.03, 3, 3])
    dl = wb.create_sheet("Daily Log")
    dl.append(["Daily Log"]); dl.append([]); dl.append([])
    dl.append(["Log ID", "Date logged", "Reviewer", "Source", "Reference no.", "TPA", "Provider", "Flag code", "Flag type",
               "Flag description", "Flag status", "Severity", "Action level", "Amount at risk (SAR)", "Owner", "Owner type",
               "Due date", "Status", "Date closed", "Resolution / outcome", "Sent to CNHI on", "CNHI response"])
    rows = [("MRS-1", 2, "R1", "H1", "EMG-01", "Confirmed", 10000), ("MRS-2", 2, "R1", "H1", "DQ-10", "Confirmed", 4000),
            ("MRS-3", 3, "R2", "H2", "DQ-10", "Confirmed", 7000), ("MRS-4", 4, "R3", "H2", "BEN-02", "Cleared on review", 500)]
    for lid, day, ref, prov, code, st, amt in rows:
        dl.append([lid, dt.date(2026, 9, day), "Dr Example", "Approval", ref, "GlobeMed", prov, code, "t",
                   "Patient narrative must never be read", st, "High", "TPA level", amt, "Named owner", "CNHI Approvals",
                   dt.date(2026, 9, 20), "Open", None, "free text", None, None])
    cf = wb.create_sheet("Config")
    cf.append(["Configuration"]); cf.append([]); cf.append([])
    cf.append(["Flag code", "Flag type (English)", "نوع الملاحظة"])
    for code in ("EMG-01", "DQ-10", "BEN-02"):
        cf.append([code, "type", "نوع"])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


class AppTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.c = A.app.test_client()
        r = cls.c.get("/login")
        r = cls.c.post("/login", data=dict(csrf=tok(r.text), email="admin@test.local", password="First!Pass2026x"))
        assert r.status_code == 302, r.status_code
        r = cls.c.get("/account")
        r = cls.c.post("/account", data=dict(csrf=tok(r.text), current="First!Pass2026x", new=NEW_PW, repeat=NEW_PW))
        assert r.status_code == 302

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TMP, ignore_errors=True)

    def upload(self, kind, content, name, label=""):
        r = self.c.get("/data")
        r = self.c.post("/data", data=dict(csrf=tok(r.text), kind=kind, label=label, file=(io.BytesIO(content), name)),
                        content_type="multipart/form-data")
        return r

    def publish(self, loc):
        r = self.c.get(loc)
        return self.c.post(loc, data=dict(csrf=tok(r.text), act="publish"))

    # ---------------------------------------------------------------- tests
    def test_01_totp(self):
        s = auth.new_mfa_secret()
        self.assertTrue(auth.check_totp(s, auth.totp(s)))
        self.assertFalse(auth.check_totp(s, "000000" if auth.totp(s) != "000000" else "111111"))

    def test_02_csrf_required(self):
        r = self.c.post("/data", data=dict(kind="F1"))
        self.assertEqual(r.status_code, 400)

    def test_03_bad_file_rejected(self):
        r = self.upload("F1", b"not a workbook", "x.xlsx")
        self.assertEqual(r.status_code, 302)
        self.assertEqual(db.val("SELECT COUNT(*) FROM uploads", default=0), 0)

    def test_04_daily_register_upload_and_kpis(self):
        r = self.upload("F1", daily_register(), "daily.xlsm")
        loc = r.headers["Location"]
        self.assertIn("/data/", loc)
        page = self.c.get(loc).text
        self.assertIn("Flags", page)
        self.publish(loc)
        data = store.current()
        flags = data["F1"]["flags"]
        # patient narrative, owner names and staff names are never stored
        blob = repr(data)
        self.assertNotIn("Patient narrative", blob)
        self.assertNotIn("Dr Example", blob)
        self.assertNotIn("Named owner", blob)
        self.assertNotEqual(flags[0]["ref"], "R1")                     # references are hashed
        self.assertEqual(data["F1"]["intake"][-1]["tpa"], "NextCare")   # TPA names normalised
        with A.app.test_request_context():
            from flask import g
            g.lang = "en"
            c = kpi.Ctx(data, "2026-09", "")
            self.assertEqual(kpi.compute("EX-01", c)["raw"], 1000)
            self.assertAlmostEqual(kpi.compute("EX-03", c)["raw"], 2 / 1000)      # 2 encounters (R1, R2) confirmed
            self.assertEqual(kpi.compute("EX-04", c)["raw"], 17000)               # R1 counted once at 10,000 (D-19)
            self.assertAlmostEqual(kpi.compute("FW-02", c)["raw"], 3 / 4)
            self.assertEqual(kpi.compute("EX-06", c)["state"], "nodata")          # column not present (D-09)
            c10 = kpi.Ctx(data, "2026-10", "")
            self.assertEqual(kpi.compute("DR-02", c10)["raw"], 10)
            self.assertAlmostEqual(kpi.compute("EX-02", c10)["raw"], 0.8)

    def test_05_disputes_template_upload(self):
        wb = load_workbook(os.path.join(ROOT, "static", "downloads", "CNHI_Hospital_Disputes_Tracker_v1.0.xlsx"))

        def put(ws, row, values):                                   # columns found by header text (row 4)
            cols = {c.value: c.column_letter for c in ws[4] if c.value}
            for k, v in values.items():
                ws[f"{cols[k]}{row}"] = v
        put(wb["Batches"], 5, {"Batch ID": "B01", "Hospital": "H1", "Received date": dt.date(2026, 9, 3),
                               "KPMG opinion date": dt.date(2026, 9, 6)})
        base = {"Batch ID": "B01", "Hospital": "H1", "Pathway": "Claim", "Dispute theme": "Coding or billing"}
        put(wb["Disputes"], 5, dict(base, **{"Dispute ID": "DSP-H1-001", "Claimed amount (SAR)": 10000,
                                             "Approved amount (SAR)": 0, "KPMG position": "Partially accepted",
                                             "Value to pay (SAR) - partial only": 6000}))
        put(wb["Disputes"], 6, dict(base, **{"Dispute ID": "DSP-H1-002", "KPMG position": "Not accepted"}))
        put(wb["Disputes"], 7, dict(base, **{"Dispute ID": "DSP-H1-003", "KPMG position": "Pending"}))
        put(wb["Actions"], 6, {"Action ID": "A-0001", "Date": dt.date(2026, 9, 3), "Dispute ID": "DSP-H1-001",
                               "Action type": "Clinical review", "Team member": "Someone", "Hours": 2.5,
                               "Estimated? (backfilled)": "Yes"})
        bio = io.BytesIO()
        wb.save(bio)                                                # saved without recalculation: the parser derives
        r = self.upload("F3", bio.getvalue(), "disputes.xlsx")
        self.publish(r.headers["Location"])
        data = store.current()
        self.assertEqual(len(data["F3"]["disputes"]), 3)                     # example action row ignored
        first = data["F3"]["disputes"][0]
        self.assertEqual((first["received"], first["value"], first["maintained"]), ("2026-09-03", 10000, 4000))
        self.assertIsNone(data["F3"]["disputes"][2]["closed"])                  # pending keeps no opinion date
        with A.app.test_request_context():
            from flask import g
            g.lang = "en"
            c = kpi.Ctx(data, "2026-09", "")
            self.assertEqual(kpi.compute("EX-07", c)["raw"], 2)
            self.assertAlmostEqual(kpi.compute("EX-08", c)["raw"], 0.5)
            self.assertEqual(kpi.compute("HD-05", c)["raw"], 4000)
            self.assertEqual(kpi.compute("HD-07", c)["raw"], 2.5)
            self.assertEqual(kpi.compute("HD-06", c)["raw"], 1)               # Thu 3 Sep -> Sun 6 Sep: Fri and Sat excluded

    def test_06_governance_template_round_trip(self):
        r = self.c.get("/data/template/governance.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.data))
        ws = wb["Deliverables"]
        ws["A6"], ws["B6"], ws["C6"], ws["E6"] = "Monthly review report", "Aug-2026", dt.date(2026, 9, 10), dt.date(2026, 9, 9)
        te = wb["Team effort"]
        for i, (w, role, days) in enumerate([("MARSAD daily review (approvals and claims)", "Physician", 44),
                                             ("Hospital disputes", "Physician", 22), ("Financial and accounting review", "Finance consultant", 22),
                                             ("Reporting and governance", "Additional KPMG support", 5)], start=6):
            te[f"A{i}"], te[f"B{i}"], te[f"C{i}"], te[f"D{i}"] = w, role, "Sep-2026", days
        rq = wb["CNHI requests"]
        rq["A6"], rq["B6"], rq["C6"], rq["D6"], rq["F6"], rq["G6"] = "Q1", dt.date(2026, 9, 1), "CNHI Finance", "Case review", dt.date(2026, 9, 3), 5
        wb["Contract"]["B5"] = dt.date(2025, 10, 1)
        bio = io.BytesIO()
        wb.save(bio)
        r = self.upload("F5", bio.getvalue(), "gov.xlsx")
        self.publish(r.headers["Location"])
        with A.app.test_request_context():
            from flask import g
            g.lang = "en"
            c = kpi.Ctx(store.current(), "2026-09", "")
            self.assertEqual(kpi.compute("EX-10", c)["raw"], 1.0)
            self.assertEqual(kpi.compute("OP-03", c)["raw"], 93)                 # example row ignored, month typed as text
            self.assertAlmostEqual(kpi.compute("OP-04", c)["raw"], 66 / 88)      # reporting left out of the mix
            self.assertEqual(kpi.compute("OP-05", c)["raw"], 27)                 # disputes 22 + additional support 5
            self.assertEqual(kpi.compute("OP-06", c)["raw"], 1)
            self.assertEqual(kpi.compute("OP-07", c)["value"], "Month 12 of 24")
            self.assertAlmostEqual(kpi.compute("OP-08", c)["raw"], (66 + 22) / 22 / 4)   # 22 working days in Sep 2026 (Sun-Thu)
        # a CNHI viewer sees the page without the KPMG-only tiles
        db.insert("users", email="viewer@test.local", name="CNHI viewer", role="cnhi", pw_hash=auth.hash_pw(NEW_PW), must_change_pw=0)
        c2 = A.app.test_client()
        r = c2.get("/login")
        c2.post("/login", data=dict(csrf=tok(r.text), email="viewer@test.local", password=NEW_PW))
        page = c2.get("/p/delivery?period=2026-09").text
        self.assertIn("OP-05", page)
        self.assertNotIn("OP-08", page)
        self.assertIn("OP-08", self.c.get("/p/delivery?period=2026-09").text)

    def test_07_every_page_both_languages(self):
        for lg in ("en", "ar"):
            self.c.get(f"/lang/{lg}")
            for p in config.PAGES:
                r = self.c.get(f"/p/{p['slug']}?period=2026-09")
                self.assertEqual(r.status_code, 200, p["slug"])
                self.assertNotIn("Could not compute", r.text, p["slug"])
        self.c.get("/lang/en")

    def test_08_period_close_freezes_figures(self):
        r = self.c.get("/periods")
        self.c.post("/periods", data=dict(csrf=tok(r.text), month="2026-09", act="close", note="test"))
        self.assertIsNotNone(db.q1("SELECT * FROM snapshots WHERE month='2026-09'"))
        r = self.c.get("/p/executive?period=2026-09")
        self.assertIn("badge ok", r.text)
        r = self.c.get("/periods")
        self.c.post("/periods", data=dict(csrf=tok(r.text), month="2026-09", act="reopen"))
        self.assertIsNone(db.q1("SELECT * FROM snapshots WHERE month='2026-09'"))

    def test_09_cnhi_viewer_cannot_administer(self):
        r = self.c.get("/admin/users")
        r = self.c.post("/admin/users", data=dict(csrf=tok(r.text), act="add", email="cfo@test.local", name="CFO", role="cnhi"))
        temp = re.search(r"Temporary password: (\S+)", r.text if r.status_code == 200 else self.c.get("/admin/users").text)
        self.assertIsNotNone(temp)
        other = A.app.test_client()
        r = other.get("/login")
        r = other.post("/login", data=dict(csrf=tok(r.text), email="cfo@test.local", password=temp.group(1)))
        self.assertEqual(r.status_code, 302)
        self.assertEqual(other.get("/data").status_code, 302)               # must change password first
        r = other.get("/account")
        other.post("/account", data=dict(csrf=tok(r.text), current=temp.group(1), new="Cfo!Pass2026xyz", repeat="Cfo!Pass2026xyz"))
        self.assertEqual(other.get("/data").status_code, 403)
        self.assertEqual(other.get("/p/executive").status_code, 200)

    def test_11_mfa_setup_with_qr_code(self):
        db.insert("users", email="mfa@test.local", name="MFA user", role="kpmg", pw_hash=auth.hash_pw(NEW_PW), must_change_pw=0)
        A.app.config["REQUIRE_MFA"] = True
        try:
            c = A.app.test_client()
            c.get("/lang/ar")
            r = c.get("/login")
            r = c.post("/login", data=dict(csrf=tok(r.text), email="mfa@test.local", password=NEW_PW))
            self.assertIn("/mfa/setup", r.headers["Location"])
            page = c.get("/mfa/setup").text
            uri = re.search(r'href="(otpauth://totp/[^"]+)"', page).group(1).replace("&amp;", "&")
            self.assertIn("issuer=MARSAD%20Insights", uri)
            try:
                import segno  # noqa: F401
                self.assertIn('class="qr"', page)                                   # QR code drawn on the page
                self.assertIn('viewBox="', page)
            except ImportError:
                self.assertIn('class="secret"', page)                              # falls back to the typed key
            secret = re.search(r"secret=([A-Z2-7]+)", uri).group(1)
            r = c.post("/mfa/setup", data=dict(csrf=tok(page), code=auth.totp(secret)))
            self.assertEqual(r.status_code, 302)
            self.assertIn("/p/executive", r.headers["Location"])
            self.assertIn('lang="ar"', c.get("/p/executive").text)                    # language chosen before sign-in kept
            self.assertEqual(db.q1("SELECT mfa_enabled FROM users WHERE email='mfa@test.local'")["mfa_enabled"], 1)
        finally:
            A.app.config["REQUIRE_MFA"] = False

    def test_10_lockout_after_failed_logins(self):
        other = A.app.test_client()
        for _ in range(auth.MAX_FAILED):
            r = other.get("/login")
            other.post("/login", data=dict(csrf=tok(r.text), email="cfo@test.local", password="wrong"))
        r = other.get("/login")
        r = other.post("/login", data=dict(csrf=tok(r.text), email="cfo@test.local", password="Cfo!Pass2026xyz"))
        self.assertEqual(r.status_code, 401)


if __name__ == "__main__":
    unittest.main()
