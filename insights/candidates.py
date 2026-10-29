"""Candidate rows for the F6 register, built from the published F1-F3 data (option 1, agreed 29 Oct 2026).

The dashboard never writes to the register. It offers an Excel file in the register's own column order; the maintainer
copies the rows CNHI should consider into F6. Each candidate carries a source key in 'Source reference'; once that key is
in F6 the candidate is not offered again.

  Decision candidates
    F1  confirmed flags at TPA level, grouped by TPA, flag code and month (not one row per encounter); per TPA and month
        only the 5 largest groups worth at least SAR 100,000 (config/settings.yaml: candidates)
    F2  GSS reports issued with an Unsatisfactory opinion (one per report)
    F3  hospital disputes with a fraud, waste or abuse indicator, grouped by hospital and indicator
  Action candidates
    F1  confirmed flags CNHI endorsed for routing to the TPA, grouped by TPA, flag code and month
Recovered values are not matched from the quarterly accounting report: it has no reference that links a deduction to an
action yet, so recoveries are entered by hand on evidence.
"""
import datetime as dt
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .xlsx_templates import F6_SHEETS

CS = config.SETTINGS.get("candidates") or {}
MIN_SAR = float(CS.get("min_value_sar", 100000))
TOP_N = int(CS.get("top_per_tpa_month", 5))

NAVY = "00338D"


def _keys(f6):
    if not f6:
        return set()
    out = set()
    for row in list(f6.get("decisions", [])) + list(f6.get("actions", [])) + list(f6.get("recommendations", [])):
        for part in (row.get("source") or "").replace(";", ",").split(","):
            if part.strip():
                out.add(part.strip().lower())
    return out


def _once(flags):
    m = {}
    for f in flags:
        k = f["ref"] or f["log_id"]
        m[k] = max(m.get(k, 0), f["amount"] or 0)
    return m


def build(data):
    """Returns (decision_rows, action_rows, skipped) where rows are dicts keyed by the F6 column headers."""
    seen = _keys(data.get("F6"))
    dec, act, skipped = [], [], 0
    below = 0

    def add(lst, row):
        nonlocal skipped
        if row["Source reference"].lower() in seen:
            skipped += 1
        else:
            lst.append(row)

    f1 = data.get("F1")
    if f1:
        tpa_level, endorsed = defaultdict(list), defaultdict(list)
        for f in f1["flags"]:
            if f["status"] != "Confirmed" or not f["tpa"] or not f["code"]:
                continue
            key = (f["tpa"], f["code"].upper(), f["date"][:7])
            if (f.get("cnhi_response") or "").lower().startswith("endorse"):
                endorsed[key].append(f)
            elif "tpa" in (f.get("action") or "").lower():
                tpa_level[key].append(f)
        ranked = defaultdict(list)
        for (tpa, code, month), fl in tpa_level.items():
            m = _once(fl)
            ranked[(tpa, month)].append((sum(m.values()), code, m))
        keep = []
        for (tpa, month), groups in ranked.items():
            groups.sort(key=lambda g: -g[0])
            for i, (v, code, m) in enumerate(groups):
                if i < TOP_N and v >= MIN_SAR:
                    keep.append((tpa, code, month, m))
                else:
                    below += 1
        for tpa, code, month, m in sorted(keep):
            add(dec, {"Subject type": "TPA", "Subject name": tpa, "Topic": "Recovery" if sum(m.values()) else "Corrective action plan",
                      "Source reference": f"F1:{tpa}:{code}:{month}", "Value at stake (SAR)": round(sum(m.values()), 2),
                      "Decision": "Pending",
                      "_basis": f"{len(m)} encounters with confirmed flag {code} in {month}, TPA level (daily register)"})
        for (tpa, code, month), fl in sorted(endorsed.items()):
            m = _once(fl)
            add(act, {"Target type": "TPA", "Target name": tpa, "TPA": tpa,
                      "Action type": "Recover overpayment" if sum(m.values()) else "Process or system fix", "Flag code": code,
                      "Source reference": f"F1E:{tpa}:{code}:{month}", "Status": "Open",
                      "Value agreed for recovery (SAR)": None,
                      "_basis": f"{len(m)} encounters endorsed by CNHI for routing to the TPA, flag {code}, {month}; "
                                f"amount at risk SAR {sum(m.values()):,.0f}. Link the Decision ID and agree the recovery value."})
    f2 = data.get("F2")
    if f2:
        at_issue = defaultdict(float)
        for s in f2["sample"]:
            at_issue[s["gss_id"]] += s.get("at_issue") or 0
        for g in f2["gss"]:
            if g.get("issued") and g.get("opinion") == "Unsatisfactory":
                add(dec, {"Subject type": "TPA", "Subject name": g["tpa"], "Topic": "Recovery",
                          "Source reference": f"F2:{g['gss_id']}", "Value at stake (SAR)": round(at_issue[g["gss_id"]], 2),
                          "Decision": "Pending",
                          "_basis": f"GSS report {g['gss_id']} issued {g['issued']} with an Unsatisfactory opinion; "
                                    f"hospital {g.get('hospital') or '-'}"})
    f3 = data.get("F3")
    if f3:
        grp = defaultdict(list)
        for d in f3["disputes"]:
            if d.get("fwa") not in (None, "", "None") and d.get("hospital"):
                grp[(d["hospital"], d["fwa"])].append(d)
        for (hosp, fwa), ds in sorted(grp.items()):
            topic = "Investigation referral" if fwa.lower().startswith("suspected") else "Recovery"
            val = sum((d.get("value") if d.get("value") is not None else d.get("claimed")) or 0 for d in ds)
            add(dec, {"Subject type": "Provider", "Subject name": hosp, "Topic": topic,
                      "Source reference": f"F3:{hosp}:{fwa}", "Value at stake (SAR)": round(val, 2) or None, "Decision": "Pending",
                      "_basis": f"{len(ds)} hospital disputes with indicator '{fwa}' ({', '.join(d['id'] for d in ds[:6])}"
                                f"{'…' if len(ds) > 6 else ''})"})
    build.below = below
    return dec, act, skipped


def to_xlsx(data):
    dec, act, skipped = build(data)
    wb = Workbook()
    gd = wb.active
    gd.title = "Guide"
    gd.column_dimensions["A"].width = 120
    lines = [
        ("F6 candidates - MARSAD Insights (KPMG internal)", True),
        (f"Built {dt.datetime.now():%Y-%m-%d %H:%M} from the published daily register, GSS workbook and disputes tracker. "
         f"{len(dec)} decision candidates and {len(act)} action candidates; {skipped} already in the register were left out.", False),
        ("Review each row. Copy the ones CNHI should consider into the matching sheet of the governance & corrective-action "
         "register (F6), give them an ID, and keep the Source reference unchanged so they are not offered again.", False),
        (f"Daily-register decision candidates are limited to the {TOP_N} largest flag codes per TPA and month worth at least "
         f"SAR {MIN_SAR:,.0f}; {getattr(build, 'below', 0)} smaller groups were left out (change the limits in config/settings.yaml).", False),
        ("The columns are in the register's order. The grey 'Candidate basis' column is for your review only; do not copy it.", False),
        ("This file can contain provider names. Keep it in the restricted engagement folder; do not share it with CNHI as is.", False),
        ("Recovered values are not suggested: the quarterly accounting report has no reference linking a deduction to an action. "
         "Enter recoveries by hand on evidence.", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = gd.cell(row=i, column=1, value=txt)
        c.font = Font(size=14 if bold else 10, bold=bold, color=NAVY if bold else None)
        c.alignment = Alignment(wrap_text=True)
    for title, sheet, rows in (("Decision candidates", "Decisions", dec), ("Action candidates", "Corrective actions", act)):
        ws = wb.create_sheet(title)
        cols = F6_SHEETS[sheet][0] + ["Candidate basis"]
        ws["A1"] = f"{title}: paste into F6 '{sheet}'"
        ws["A1"].font = Font(size=13, bold=True, color=NAVY)
        for j, h in enumerate(cols, start=1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="808080" if h == "Candidate basis" else NAVY)
            ws.column_dimensions[get_column_letter(j)].width = 60 if h == "Candidate basis" else 18
        for i, r in enumerate(rows, start=5):
            for j, h in enumerate(cols, start=1):
                v = r.get("_basis") if h == "Candidate basis" else r.get(h)
                c = ws.cell(row=i, column=j, value=v)
                if h == "Candidate basis":
                    c.font = Font(italic=True, color="595959")
                    c.alignment = Alignment(wrap_text=True)
                if "SAR" in h and v is not None:
                    c.number_format = "#,##0"
        ws.freeze_panes = "A5"
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), len(dec), len(act), skipped
