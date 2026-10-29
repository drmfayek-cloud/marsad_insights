"""Builds the governance register template (F5) offered for download on the Data uploads page."""
import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

F, NAVY = "Arial", "00338D"
thin = Side(style="thin", color="C8CED9")
B = Border(left=thin, right=thin, top=thin, bottom=thin)

WORKSTREAMS = ["MARSAD daily review (approvals and claims)", "GSS claims assurance", "Hospital disputes",
               "Financial and accounting review", "Reporting and governance", "CNHI ad-hoc requests"]
ROLES = ["Physician", "Finance consultant", "Additional KPMG support"]
REQUEST_TYPES = ["Data extract or analysis", "Case review", "Report or presentation", "Meeting or workshop", "Other"]

SHEETS = {
    "Deliverables": (["Deliverable type", "Period", "Due date", "QA sign-off date", "Submitted date", "CNHI approval date", "Reference"],
                     [22, 14, 13, 14, 13, 15, 24],
                     ["EXAMPLE Monthly review report", "Sep-2026", dt.date(2026, 10, 10), dt.date(2026, 10, 8),
                      dt.date(2026, 10, 9), dt.date(2026, 10, 20), "MRR-2026-09"],
                     {"A": ["Monthly review report", "Quarterly financial report", "Annual consolidated report",
                            "GSS assurance report"]}),
    "Meetings": (["Meeting type", "Planned date", "Held date", "Note"], [26, 13, 13, 40],
                 ["EXAMPLE Weekly project meeting", dt.date(2026, 10, 5), dt.date(2026, 10, 5), ""],
                 {"A": ["Weekly project meeting", "Monthly management review", "Daily coordination"]}),
    "Recommendations": (["Recommendation ID", "Title", "Root cause", "Owner", "Raised date", "Status", "Accepted date",
                         "Implemented date"], [16, 40, 24, 20, 13, 14, 13, 15],
                        ["EXAMPLE-R01", "Require bed-charge line on inpatient GSS claims", "TPA adjudication", "CNHI Claims",
                         dt.date(2026, 10, 1), "Raised", None, None],
                        {"F": ["Raised", "Accepted", "Implemented", "Rejected"]}),
    "Knowledge transfer": (["Item", "Type", "Planned date", "Delivered date", "Audience"], [40, 16, 13, 14, 24],
                           ["EXAMPLE Review methodology walkthrough", "Session", dt.date(2026, 11, 15), None, "CNHI PFPS team"],
                           {"B": ["Session", "Document", "Tool handover"]}),
    "Team effort": (["Workstream", "Role", "Month", "Person-days", "Note"], [44, 26, 13, 13, 40],
                    ["EXAMPLE MARSAD daily review (approvals and claims)", "Physician", dt.date(2026, 10, 1), 40,
                     "One row per workstream and role each month"],
                    {"A": WORKSTREAMS, "B": ROLES}),
    "CNHI requests": (["Request ID", "Date received", "Requesting unit", "Request type", "Due date", "Delivered date",
                       "Hours", "Status", "Note"], [14, 14, 26, 26, 13, 14, 9, 13, 40],
                      ["EXAMPLE-Q01", dt.date(2026, 10, 5), "CNHI Finance", "Data extract or analysis", dt.date(2026, 10, 8),
                       dt.date(2026, 10, 7), 6, "Delivered", "Unit names only, no personal names"],
                      {"D": REQUEST_TYPES, "H": ["Open", "Delivered", "Cancelled"]}),
}

# Contract baseline (sheet 'Contract'): item -> (default, note)
CONTRACT = [
    ("Contract start date", None, "First day of the 24-month engagement. Needed for contract-to-date totals and month n of 24."),
    ("Contract duration (months)", 24, ""),
    ("Physicians (FTE)", 3, "Contracted physician reviewers: pre-authorizations and claims"),
    ("Finance consultants (FTE)", 1, "Contracted finance consultant: financial and accounting review"),
    ("Hours per person-day", 8, "Converts dispute and request hours into person-days"),
]


def governance_template():
    wb = Workbook()
    gd = wb.active
    gd.title = "Guide"
    gd.column_dimensions["A"].width = 110
    lines = [
        ("Governance register - MARSAD Insights (F5)", True),
        ("Keep this file up to date and upload it under Data uploads. It feeds EX-10, GV-01 to GV-06, FW-08 and the 'KPMG delivery & effort' page (OP-01 to OP-09).", False),
        ("Row 5 of each sheet is an example (its first cell starts with EXAMPLE): it is ignored on upload; replace it.", False),
        ("Deliverables: one row per deliverable and period. A deliverable is on time when Submitted date <= Due date.", False),
        ("Meetings: one row per planned meeting; enter the Held date when it takes place.", False),
        ("Recommendations: process-improvement recommendations to CNHI and their status.", False),
        ("Knowledge transfer: sessions and documents planned for CNHI, with delivery dates.", False),
        ("Team effort: each month, the person-days spent per workstream and role (monthly totals, no timesheets). "
         "'Additional KPMG support' is effort from outside the contracted team (management, QA, specialists).", False),
        ("CNHI requests: ad-hoc requests from CNHI outside the routine deliverables. Record the requesting unit, never a person's name. "
         "The Note column is not read by the dashboard.", False),
        ("Contract: the contract baseline (start date, duration, contracted team). Enter the start date once.", False),
        ("Dates: enter real Excel dates. Do not insert rows above the header row (row 4).", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = gd.cell(row=i, column=1, value=txt)
        c.font = Font(name=F, size=14 if bold else 10, bold=bold, color=NAVY if bold else None)
        c.alignment = Alignment(wrap_text=True)
    for name, (cols, widths, example, lists) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws["A1"] = name
        ws["A1"].font = Font(name=F, size=13, bold=True, color=NAVY)
        ws["A2"] = "Blue text: you enter. Row 5 is an example."
        ws["A2"].font = Font(name=F, size=9, italic=True)
        for j, (h, w) in enumerate(zip(cols, widths), start=1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.border = B
            ws.column_dimensions[get_column_letter(j)].width = w
        for j, v in enumerate(example, start=1):
            c = ws.cell(row=5, column=j, value=v)
            c.font = Font(name=F, size=10, italic=True, color="808080")
            if isinstance(v, dt.date):
                c.number_format = "dd-mmm-yyyy"
        for r in range(6, 506):
            for j, h in enumerate(cols, start=1):
                c = ws.cell(row=r, column=j)
                c.font = Font(name=F, size=10, color="0000FF")
                c.border = B
                if "date" in h.lower():
                    c.number_format = "dd-mmm-yyyy"
        for col, vals in lists.items():
            dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}5:{col}505")
        ws.freeze_panes = "A5"
    ws = wb.create_sheet("Contract")
    ws["A1"] = "Contract"
    ws["A1"].font = Font(name=F, size=13, bold=True, color=NAVY)
    ws["A2"] = "Contract baseline used by the 'KPMG delivery & effort' page. Yellow = to complete."
    ws["A2"].font = Font(name=F, size=9, italic=True)
    for j, (h, w) in enumerate(zip(["Item", "Value", "Note"], [30, 16, 80]), start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name=F, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.border = B
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, (item, val, note) in enumerate(CONTRACT, start=5):
        ws.cell(row=i, column=1, value=item).font = Font(name=F, size=10, bold=True)
        v = ws.cell(row=i, column=2, value=val)
        v.font = Font(name=F, size=10, color="0000FF")
        if "date" in item.lower():
            v.number_format = "dd-mmm-yyyy"
        if val is None:
            v.fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(row=i, column=3, value=note).font = Font(name=F, size=9, italic=True)
        for j in (1, 2, 3):
            ws.cell(row=i, column=j).border = B
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
